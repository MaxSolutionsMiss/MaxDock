"""MaxDock deck kit — the design system of the MaxDock introduction deck,
expressed as python-pptx primitives so new slides match the existing ones.

Every measurement is in inches on a 13.333 x 7.5 canvas. The vocabulary is
deliberately small: a slide frame (eyebrow, headline, deck line, footer), cards,
chips, rules, bars and text. Anything a new slide needs that is not here should
be added here rather than drawn inline, so the next deck revision stays cheap.
"""
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------- palette ---
INK = "0B1F33"        # headlines, dark panels
DEEP = "072238"       # darker panel wash
BODY = "456174"       # body copy
MUTED = "6F8491"      # secondary copy, footer
FAINT = "7FA3B8"      # captions on dark
WHITE = "FFFFFF"
BG = "F5F8FA"         # slide background
LINE = "D8E4EA"       # hairlines and card borders
PANEL = "E7EDF1"      # inert fill
PANEL_SOFT = "EDF2F6"
BLUE = "0877C9"
CYAN = "14AEE8"
BLUE_TINT = "E4EFF9"
GREEN = "14A66A"
GREEN_TINT = "E1F3EA"
GREEN_DARK = "10603F"
AMBER = "F2A11B"
AMBER_TINT = "FCF0DC"
AMBER_DARK = "8A5B06"
RED = "D84C4C"
RED_TINT = "FAE7E7"
RED_DARK = "8E2F2F"
ICE = "B9D7E8"        # body copy on a navy panel
NAVY_RULE = "1E4A63"  # hairline on a navy panel

DISPLAY = "Aptos Display"   # headlines and card titles
TEXT = "Aptos"              # everything else

# ------------------------------------------------------------------ frame ---
MARGIN_L = 0.58
CONTENT_W = 12.17
FOOT_Y = 7.07


def _rgb(hexstr):
    return RGBColor.from_string(hexstr)


def _emu(v):
    return Emu(int(round(v * 914400)))


def text(slide, x, y, w, h, runs, size=11.5, color=BODY, bold=False,
         italic=False, font=TEXT, align="l", anchor="t", line_spacing=1.0,
         space_after=0, spc=None, wrap=True):
    """A text box with no internal padding, so x/y are the true text corner.

    `runs` is a string, a list of strings (one paragraph each), or a list of
    lists of (text, overrides) tuples for mixed formatting within a paragraph.
    """
    box = slide.shapes.add_textbox(_emu(x), _emu(y), _emu(w), _emu(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE,
                          "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(runs, str):
        runs = [runs]

    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER,
                       "r": PP_ALIGN.RIGHT}[align]
        p.line_spacing = line_spacing
        if space_after:
            p.space_after = Pt(space_after)
        pieces = para if isinstance(para, list) else [(para, {})]
        for chunk, over in pieces:
            r = p.add_run()
            r.text = chunk
            f = r.font
            f.name = over.get("font", font)
            f.size = Pt(over.get("size", size))
            f.bold = over.get("bold", bold)
            f.italic = over.get("italic", italic)
            f.color.rgb = _rgb(over.get("color", color))
            letter_spacing = over.get("spc", spc)
            if letter_spacing:
                r.font._rPr.set("spc", str(int(letter_spacing)))
    return box


def rect(slide, x, y, w, h, fill=WHITE, line=None, radius=None, line_w=0.75):
    """A rectangle. `radius` is a corner radius in inches, not a ratio, so a
    card and a chip drawn side by side round by the same amount."""
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(shape_type, _emu(x), _emu(y), _emu(w), _emu(h))
    if radius:
        short = max(min(w, h), 0.01)
        sh.adjustments[0] = min(radius / short, 0.5)
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = _rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = _rgb(line)
        sh.line.width = Pt(line_w)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    return sh


CARD_RADIUS = 0.073
CHIP_RADIUS = 0.06


def card(slide, x, y, w, h, fill=WHITE, line=LINE):
    """The deck's standard card: white, hairline border, gently rounded."""
    return rect(slide, x, y, w, h, fill=fill, line=line, radius=CARD_RADIUS)


def chip(slide, x, y, w, h, label, fill=BLUE_TINT, color=BLUE, size=9.5,
         bold=True, align="c", font=TEXT, spc=None, radius=CHIP_RADIUS):
    """A small rounded label: status pills, time stamps, figure callouts."""
    rect(slide, x, y, w, h, fill=fill, line=None, radius=radius)
    text(slide, x, y, w, h, label, size=size, color=color, bold=bold,
         align=align, anchor="m", font=font, spc=spc)


def rule(slide, x, y, w, color=LINE, weight=0.75):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, _emu(x), _emu(y),
                                _emu(w), _emu(weight / 72.0))
    sh.fill.solid()
    sh.fill.fore_color.rgb = _rgb(color)
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def vrule(slide, x, y, h, color=LINE, weight=0.75):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, _emu(x), _emu(y),
                                _emu(weight / 72.0), _emu(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = _rgb(color)
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def bar(slide, x, y, w, h, pct, fill=BLUE, track=PANEL, radius=0.03):
    """A capacity bar. `pct` is 0-1 of the track width."""
    rect(slide, x, y, w, h, fill=track, line=None, radius=radius)
    if pct > 0:
        rect(slide, x, y, max(w * min(pct, 1.0), h), h, fill=fill, line=None,
             radius=radius)


def numbered_dot(slide, x, y, d, label, fill=BLUE, color=WHITE, size=10):
    rect(slide, x, y, d, d, fill=fill, line=None, radius=CHIP_RADIUS)
    text(slide, x, y, d, d, label, size=size, color=color, bold=True,
         align="c", anchor="m")


# ------------------------------------------------------------ slide frame ---
def new_slide(prs, eyebrow, headline, deck_line, logo_png):
    """A blank MaxDock slide: background, footer furniture and the header
    block every content slide shares."""
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    _set_bg(slide, BG)
    return frame(slide, eyebrow, headline, deck_line, logo_png)


def clear(slide):
    """Strip every shape off a slide so it can be redrawn in place. Redrawing
    beats delete-and-add: the slide keeps its position and its part name."""
    tree = slide.shapes._spTree
    for shape in list(slide.shapes):
        tree.remove(shape._element)
    return slide


def frame(slide, eyebrow, headline, deck_line, logo_png):
    """Draw the shared furniture — footer, eyebrow, headline, deck line."""
    rule(slide, 0.55, FOOT_Y, 12.24, LINE)
    slide.shapes.add_picture(logo_png, _emu(0.55), _emu(7.16),
                             _emu(0.16), _emu(0.16))
    text(slide, 0.78, 7.16, 1.10, 0.16, "MaxDock", size=6.7, color=MUTED)
    text(slide, 6.42, 7.16, 0.50, 0.16, "0", size=6.7, color=MUTED, align="c")
    text(slide, 11.20, 7.16, 1.58, 0.16, "MAX SOLUTIONS", size=6.7, color=MUTED,
         bold=True, align="r", spc=60)

    text(slide, MARGIN_L, 0.34, CONTENT_W, 0.22, eyebrow, size=9, color=BLUE,
         bold=True, spc=90)
    text(slide, MARGIN_L, 0.58, CONTENT_W, 0.62, headline, size=25, color=INK,
         bold=True, font=DISPLAY, line_spacing=1.02)
    if deck_line:
        text(slide, MARGIN_L, 1.26, CONTENT_W, 0.30, deck_line, size=12,
             color=BODY, line_spacing=1.1)
    return slide


def _set_bg(slide, hexstr):
    from lxml import etree
    ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    bg = etree.Element(f"{{{ns}}}bg")
    bgpr = etree.SubElement(bg, f"{{{ns}}}bgPr")
    fill = etree.SubElement(bgpr, f"{{{a}}}solidFill")
    etree.SubElement(fill, f"{{{a}}}srgbClr").set("val", hexstr)
    etree.SubElement(bgpr, f"{{{a}}}effectLst")
    slide._element.find(f"{{{ns}}}cSld").insert(0, bg)


def find(slide, x, y, tol=0.02):
    """The shape whose top-left corner sits at (x, y), in inches."""
    for sh in slide.shapes:
        if sh.left is None:
            continue
        if abs(sh.left / 914400 - x) < tol and abs(sh.top / 914400 - y) < tol:
            return sh
    raise LookupError(f"no shape at {x}, {y}")


def move(shape, x=None, y=None, w=None, h=None):
    """Reposition or resize an existing shape, in inches."""
    if x is not None:
        shape.left = _emu(x)
    if y is not None:
        shape.top = _emu(y)
    if w is not None:
        shape.width = _emu(w)
    if h is not None:
        shape.height = _emu(h)
    return shape


def set_text(shape, text):
    """Replace a shape's words, keeping the formatting of its first run."""
    tf = shape.text_frame
    for para in list(tf.paragraphs)[1:]:
        para._p.getparent().remove(para._p)
    para = tf.paragraphs[0]
    runs = para.runs
    if not runs:
        raise ValueError("shape has no run to inherit formatting from")
    runs[0].text = text
    for extra in runs[1:]:
        extra._r.getparent().remove(extra._r)
    return shape


def members(slide, anchor):
    """Every shape standing on `anchor` — a card and its contents move as one.

    Membership is by centre point, so a hairline drawn flush to a card edge
    still counts, and a neighbouring card never does.
    """
    x0, y0 = anchor.left, anchor.top
    x1, y1 = x0 + anchor.width, y0 + anchor.height
    out = []
    for sh in slide.shapes:
        # identity by element: python-pptx hands out a fresh proxy per visit,
        # so `sh is anchor` would let the anchor move itself twice
        if sh._element is anchor._element or sh.left is None:
            continue
        cx, cy = sh.left + sh.width / 2, sh.top + sh.height / 2
        if x0 - 1000 <= cx <= x1 + 1000 and y0 - 1000 <= cy <= y1 + 1000:
            out.append(sh)
    return out


def move_block(slide, x_old, y_old, x_new, y_new, w=None, h=None, text_dw=0.0):
    """Move a card (or band) and everything on it, optionally resizing.

    `text_dw` is applied to the width of any text box wide enough to be body
    copy, so narrowing a card reflows its paragraphs instead of letting them
    run over the edge. Full-height children — the text inside a banner — track
    the new height so they stay centred.
    """
    anchor = find(slide, x_old, y_old)
    kids = members(slide, anchor)
    old_h = anchor.height
    dx, dy = _emu(x_new - x_old), _emu(y_new - y_old)

    anchor.left, anchor.top = _emu(x_new), _emu(y_new)
    if w is not None:
        anchor.width = _emu(w)
    if h is not None:
        anchor.height = _emu(h)

    for sh in kids:
        sh.left += dx
        sh.top += dy
        if text_dw and sh.has_text_frame and sh.width > anchor.width * 0.4:
            sh.width = max(_emu(0.3), sh.width + _emu(text_dw))
        if h is not None and sh.has_text_frame and sh.height >= old_h - _emu(0.2):
            sh.height = anchor.height
    return anchor


def shift_band(slide, y_from, y_to, dy, x_from=None, x_to=None):
    """Nudge every shape whose top edge sits in a horizontal band.

    Pass x_from/x_to to limit the nudge to one column of the slide, so a chart
    on the left can move without dragging the one on the right with it.
    """
    lo, hi = _emu(y_from), _emu(y_to)
    xlo = _emu(x_from) if x_from is not None else None
    xhi = _emu(x_to) if x_to is not None else None
    for sh in slide.shapes:
        if sh.top is None or not (lo <= sh.top <= hi):
            continue
        if xlo is not None and sh.left < xlo:
            continue
        if xhi is not None and sh.left > xhi:
            continue
        sh.top += _emu(dy)


def renumber(prs):
    """Rewrite the footer page number on every slide to its position."""
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if not sh.has_text_frame or sh.left is None:
                continue
            if abs(sh.left / 914400 - 6.42) < 0.02 and abs(sh.top / 914400 - 7.16) < 0.02:
                for p in sh.text_frame.paragraphs:
                    for j, r in enumerate(p.runs):
                        r.text = str(i) if j == 0 else ""


def place(prs, slide, index):
    """Move `slide` to a 0-based position, found by its own relationship id."""
    rid_attr = ("{http://schemas.openxmlformats.org/officeDocument/2006/"
                "relationships}id")
    target = None
    for sld_id in prs.slides._sldIdLst:
        if prs.part.related_part(sld_id.get(rid_attr)) is slide.part:
            target = sld_id
            break
    if target is None:
        raise ValueError("slide is not in this presentation")
    lst = prs.slides._sldIdLst
    lst.remove(target)
    lst.insert(index, target)
